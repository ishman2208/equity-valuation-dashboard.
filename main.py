"""
Financial Valuation Dashboard
=============================
Enter a ticker (e.g. AAPL, MSFT, BHP.AX) and get an Excel workbook containing
a DCF, Comparable Company Analysis, historical financials, a sensitivity
table and a BUY / SELL recommendation.

Run locally :  python main.py
Run in Colab:  paste this whole file into one cell and run it (or upload it and
               use  %run main.py ). The finished .xlsx downloads automatically.

Simplifications (deliberate, for transparency):
  * Revenue grows at the user's terminal growth rate in every forecast year.
  * FCFF ratios (EBIT margin, D&A %, CapEx %) use recent historical averages.
  * Cash flows are discounted at year-end (no mid-year convention).
  * Recommendation is BUY if upside >= 15%, otherwise SELL.
"""

import datetime as dt
import math
import os
import subprocess
import sys
import warnings

warnings.filterwarnings("ignore")

IN_COLAB = "google.colab" in sys.modules


# --------------------------------------------------------------------------
# 0. Dependencies (auto-install, mainly for Google Colab)
# --------------------------------------------------------------------------
def ensure_packages():
    needed = {"yfinance": "yfinance", "pandas": "pandas", "numpy": "numpy", "openpyxl": "openpyxl"}
    missing = []
    for module, pip_name in needed.items():
        try:
            __import__(module)
        except ImportError:
            missing.append(pip_name)
    if IN_COLAB:
        missing.append("yfinance")  # Colab's bundled yfinance is often too old for Yahoo's API
    if missing:
        print("Installing required packages:", ", ".join(sorted(set(missing))))
        args = [sys.executable, "-m", "pip", "install", "-q"]
        if IN_COLAB:
            args.append("-U")
        subprocess.check_call(args + sorted(set(missing)))


ensure_packages()

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import yfinance as yf  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.series import DataPoint  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------
FONT = "Calibri"
NAVY, BLUE, LIGHT, GREY = "1F3864", "2F5597", "D9E2F3", "F2F2F2"
GREEN_BG, GREEN_FG, RED_BG, RED_FG = "C6EFCE", "006100", "FFC7CE", "9C0006"
INPUT_BLUE = "0000FF"

FMT_M = '#,##0.0;(#,##0.0);"-"'
FMT_PCT = "0.0%"
FMT_X = '0.0"x"'
FMT_NUM = "#,##0.00"

MIN_UPSIDE = 0.15
FORECAST_YEARS = 5
DEFAULTS = {"g": 0.03, "wacc": 0.09, "exit": 11.0, "w_dcf": 0.60, "w_cca": 0.40}
RANGES = {"g": (0.02, 0.04), "wacc": (0.07, 0.12), "exit": (8.0, 15.0)}
ERP = 0.05                 # equity risk premium (assumption)
FALLBACK_RF = 0.04         # used if the 10Y yield cannot be downloaded
FALLBACK_TAX = 0.25
SENS_WACC = [0.07, 0.08, 0.09, 0.10, 0.11, 0.12]
SENS_G = [0.02, 0.025, 0.03, 0.035, 0.04]
CURRENCY_SYMBOLS = {"USD": "$", "AUD": "A$", "GBP": "£", "EUR": "€", "JPY": "¥", "CAD": "C$", "HKD": "HK$", "INR": "₹"}


class DataError(Exception):
    """Raised when the ticker or essential data cannot be retrieved."""


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def num(x):
    """Convert anything to float, returning NaN if not possible."""
    try:
        v = float(x)
        return v if math.isfinite(v) else float("nan")
    except (TypeError, ValueError):
        return float("nan")


def ok(x):
    return isinstance(x, (int, float, np.floating)) and math.isfinite(x)


def pick(df, names):
    """First matching row of a yfinance statement as a numeric Series."""
    if df is None or getattr(df, "empty", True):
        return pd.Series(dtype=float)
    for n in names:
        if n in df.index:
            return pd.to_numeric(df.loc[n], errors="coerce")
    return pd.Series(np.nan, index=df.columns)


def by_year(series):
    """Re-index a statement row by fiscal year (int)."""
    s = series.copy()
    s.index = [c.year for c in pd.to_datetime(s.index)]
    return s[~s.index.duplicated()]


def safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


# --------------------------------------------------------------------------
# 1. Data retrieval
# --------------------------------------------------------------------------
def get_stock_price(ticker_obj):
    """Return (current price, 1y closing-price Series)."""
    hist = safe(lambda: ticker_obj.history(period="1y")["Close"].dropna(), pd.Series(dtype=float))
    info = safe(lambda: ticker_obj.info, {}) or {}
    price = num(info.get("currentPrice"))
    if not ok(price):
        price = num(info.get("regularMarketPrice"))
    if not ok(price) and len(hist):
        price = float(hist.iloc[-1])
    return price, hist


def get_historical_financials(ticker_obj):
    """Annual statements (oldest -> newest) merged into one DataFrame, in raw currency units."""
    inc = safe(lambda: ticker_obj.financials)
    bal = safe(lambda: ticker_obj.balance_sheet)
    cf = safe(lambda: ticker_obj.cashflow)
    if inc is None or inc.empty:
        raise DataError("No annual income-statement data is available for this ticker.")

    rev = by_year(pick(inc, ["Total Revenue", "Operating Revenue"]))
    years = sorted(rev.dropna().index)[-5:]
    if not years:
        raise DataError("Revenue history is not available for this ticker.")

    def get(df, names):
        return by_year(pick(df, names)).reindex(years)

    h = pd.DataFrame(index=years)
    h["Revenue"] = rev.reindex(years)
    h["Gross Profit"] = get(inc, ["Gross Profit"])
    ebit_op = get(inc, ["Operating Income", "EBIT"])
    h["EBIT"] = ebit_op
    h["Net Income"] = get(inc, ["Net Income", "Net Income Common Stockholders"])
    h["D&A"] = get(cf, ["Depreciation And Amortization", "Depreciation Amortization Depletion",
                        "Depreciation & amortization"]).fillna(get(inc, ["Reconciled Depreciation"]))
    h["EBITDA"] = get(inc, ["EBITDA", "Normalized EBITDA"]).fillna(h["EBIT"] + h["D&A"])
    h["CapEx"] = get(cf, ["Capital Expenditure", "Purchase Of PPE"]).abs()
    h["FCF"] = get(cf, ["Free Cash Flow"])
    h["Total Debt"] = get(bal, ["Total Debt"])
    h["Cash"] = get(bal, ["Cash Cash Equivalents And Short Term Investments", "Cash And Cash Equivalents"])
    h["Equity"] = get(bal, ["Stockholders Equity", "Common Stock Equity"])
    h["Shares"] = get(bal, ["Ordinary Shares Number", "Share Issued"])
    ca = get(bal, ["Current Assets"])
    cl = get(bal, ["Current Liabilities"])
    cdebt = get(bal, ["Current Debt", "Current Debt And Capital Lease Obligation"]).fillna(0)
    nwc = (ca - h["Cash"].fillna(0)) - (cl - cdebt)  # operating NWC: excludes cash and short-term debt
    h["NWC"] = nwc.fillna(get(bal, ["Working Capital"]))
    h["Pretax"] = get(inc, ["Pretax Income"])
    h["Tax"] = get(inc, ["Tax Provision"])
    h["Interest"] = get(inc, ["Interest Expense"]).abs()

    h["Rev Growth"] = h["Revenue"].pct_change()
    h["EBIT Margin"] = h["EBIT"] / h["Revenue"]
    h["Gross Margin"] = h["Gross Profit"] / h["Revenue"]
    h["Net Margin"] = h["Net Income"] / h["Revenue"]
    h["Debt/Equity"] = h["Total Debt"] / h["Equity"]
    return h


def get_risk_free_rate():
    """10Y yield from Yahoo (^TNX). Returns (rate, source label)."""
    v = safe(lambda: yf.Ticker("^TNX").history(period="5d")["Close"].dropna().iloc[-1])
    if v is not None and 0 < v < 20:
        return float(v) / 100, "Market data: US 10Y Treasury yield (^TNX)"
    return FALLBACK_RF, "Assumption: 10Y yield could not be downloaded"


def get_fx_rate(fin_ccy, px_ccy):
    if not fin_ccy or not px_ccy or fin_ccy == px_ccy:
        return 1.0, None
    v = safe(lambda: yf.Ticker(f"{fin_ccy}{px_ccy}=X").history(period="5d")["Close"].dropna().iloc[-1])
    if v is not None and v > 0:
        return float(v), f"Financials are in {fin_ccy}; converted to {px_ccy} at {v:.4f}."
    return 1.0, f"Financials are in {fin_ccy} but price is in {px_ccy}; FX rate unavailable, so NO conversion was applied."


def get_company_data(ticker):
    """Everything about the target company, with notes on any assumption used."""
    t = yf.Ticker(ticker)
    info = safe(lambda: t.info, {}) or {}
    price, prices = get_stock_price(t)
    if not ok(price):
        raise DataError(f"Could not find a share price for '{ticker}'. Check the ticker (e.g. AAPL or BHP.AX).")

    hist = get_historical_financials(t)
    notes, flags = [], {}
    last = lambda col: (hist[col].dropna().iloc[-1] if hist[col].notna().any() else float("nan"))

    px_ccy = info.get("currency") or "USD"
    fin_ccy = info.get("financialCurrency") or px_ccy
    fx, fx_note = get_fx_rate(fin_ccy, px_ccy)
    if fx_note:
        notes.append(fx_note)

    shares = num(info.get("sharesOutstanding"))
    if not ok(shares):
        mc = num(info.get("marketCap"))
        shares = mc / price if ok(mc) else last("Shares")
        notes.append("Assumption: shares outstanding derived from market cap / price or balance sheet.")
    if not ok(shares) or shares <= 0:
        raise DataError("Shares outstanding are not available, so a per-share value cannot be calculated.")

    market_cap = num(info.get("marketCap"))
    if not ok(market_cap):
        market_cap = price * shares
        notes.append("Assumption: market cap calculated as price x shares outstanding.")

    debt, cash = last("Total Debt"), last("Cash")
    if not ok(debt):
        debt = 0.0
        notes.append("Assumption: total debt not reported - assumed 0.")
    if not ok(cash):
        cash = 0.0
        notes.append("Assumption: cash not reported - assumed 0.")

    ev = num(info.get("enterpriseValue"))
    if not ok(ev):
        ev = market_cap + (debt - cash) * fx
        notes.append("Assumption: enterprise value calculated as market cap + net debt.")

    ebitda_ltm = num(info.get("ebitda"))
    if not ok(ebitda_ltm):
        ebitda_ltm = last("EBITDA")
    pe_now = num(info.get("trailingPE"))
    # Derive EPS from price / P/E so it is always in the same currency as the share price
    eps = price / pe_now if ok(pe_now) and pe_now > 0 else num(info.get("trailingEps"))

    return {
        "ticker": ticker, "name": info.get("longName") or info.get("shortName") or ticker,
        "sector": info.get("sector") or "N/A", "industry": info.get("industry") or "N/A",
        "industry_key": info.get("industryKey"), "sector_key": info.get("sectorKey"),
        "currency": px_ccy, "fin_currency": fin_ccy, "fx": fx,
        "price": price, "prices": prices, "market_cap": market_cap, "ev": ev, "shares": shares,
        "beta": num(info.get("beta")), "eps": eps, "ebitda": ebitda_ltm,
        "debt": debt, "cash": cash, "net_debt": debt - cash, "hist": hist, "notes": notes,
    }


# --------------------------------------------------------------------------
# 2. WACC and forecast assumptions
# --------------------------------------------------------------------------
def calculate_wacc(co):
    """Return a dict of WACC inputs (each with a source label) and the resulting WACC."""
    h = co["hist"]
    rf, rf_src = get_risk_free_rate()
    beta = co["beta"]
    beta_src = "Market data: Yahoo Finance beta"
    if not ok(beta):
        beta, beta_src = 1.0, "Assumption: beta unavailable, market beta of 1.0 used"
    coe = rf + beta * ERP

    debt = co["debt"]
    interest = h["Interest"].dropna().iloc[-1] if h["Interest"].notna().any() else float("nan")
    if debt > 0 and ok(interest) and 0.01 <= interest / debt <= 0.15:
        kd, kd_src = interest / debt, "Historical data: interest expense / total debt"
    else:
        kd, kd_src = rf + 0.02, "Assumption: risk-free rate + 2.0% credit spread"

    tax, tax_src = effective_tax_rate(co)
    equity = co["market_cap"]
    d, e = debt * co["fx"], equity
    v = d + e
    we, wd = e / v, d / v
    wacc = we * coe + wd * kd * (1 - tax)
    return {
        "rf": (rf, rf_src), "beta": (beta, beta_src), "erp": (ERP, "Assumption: equity risk premium"),
        "coe": (coe, "Calculated: Rf + Beta x ERP"), "kd": (kd, kd_src), "tax": (tax, tax_src),
        "E": (equity, "Market data: market capitalisation"),
        "D": (d, "Historical data: book value of debt used as proxy for market value"),
        "we": (we, "Calculated: E / (D+E)"), "wd": (wd, "Calculated: D / (D+E)"),
        "wacc": wacc,
    }


def effective_tax_rate(co):
    h = co["hist"]
    rates = (h["Tax"] / h["Pretax"]).replace([np.inf, -np.inf], np.nan).dropna()
    rates = rates[(rates > 0.05) & (rates < 0.45)]
    if len(rates):
        return float(rates.iloc[-1]), "Historical data: latest effective tax rate"
    default = 0.30 if co["ticker"].upper().endswith(".AX") else FALLBACK_TAX
    return default, "Assumption: effective tax rate unavailable, statutory-style rate used"


def build_forecast_assumptions(co, growth):
    """Historical-based drivers for the FCFF projection, each with a source label."""
    h = co["hist"]
    rev_pos = h["Revenue"].replace(0, np.nan)
    src = "Historical data"

    def avg_ratio(col, n=3):
        r = (h[col] / rev_pos).dropna().tail(n)
        return float(r.mean()) if len(r) else float("nan")

    a, tags = {}, {}
    a["base_revenue"], tags["base_revenue"] = h["Revenue"].dropna().iloc[-1], f"{src}: latest annual revenue"
    a["growth"], tags["growth"] = growth, "User assumption: terminal growth rate applied to all years"
    m = avg_ratio("EBIT")
    a["ebit_margin"], tags["ebit_margin"] = (m, f"{src}: 3-yr average EBIT margin") if ok(m) else (0.10, "Assumption: EBIT unavailable, 10% margin used")
    m = avg_ratio("D&A")
    a["da_pct"], tags["da_pct"] = (m, f"{src}: 3-yr average D&A % of revenue") if ok(m) else (0.04, "Assumption: D&A unavailable, 4% of revenue used")
    m = avg_ratio("CapEx")
    a["capex_pct"], tags["capex_pct"] = (m, f"{src}: 3-yr average CapEx % of revenue") if ok(m) else (0.04, "Assumption: CapEx unavailable, 4% of revenue used")
    nwc_last = h["NWC"].dropna()
    if len(nwc_last):
        a["nwc_pct"], tags["nwc_pct"] = float(nwc_last.iloc[-1] / h["Revenue"].loc[nwc_last.index[-1]]), f"{src}: latest NWC % of revenue"
    else:
        a["nwc_pct"], tags["nwc_pct"] = 0.0, "Assumption: working capital unavailable, no NWC investment assumed"
    a["tax"], tags["tax"] = effective_tax_rate(co)
    return a, tags


# --------------------------------------------------------------------------
# 3. DCF
# --------------------------------------------------------------------------
def project_fcff(a, growth=None):
    """Five-year FCFF projection. Returns a DataFrame indexed by row name, columns Year 1..5."""
    g = a["growth"] if growth is None else growth
    rev_prev, cols = a["base_revenue"], {}
    for y in range(1, FORECAST_YEARS + 1):
        rev = rev_prev * (1 + g)
        ebit = rev * a["ebit_margin"]
        tax = ebit * a["tax"]
        nopat = ebit - tax
        da, capex = rev * a["da_pct"], rev * a["capex_pct"]
        dnwc = a["nwc_pct"] * (rev - rev_prev)
        cols[f"Year {y}"] = {
            "Revenue": rev, "Revenue Growth": g, "EBIT": ebit, "EBIT Margin": a["ebit_margin"],
            "Taxes on EBIT": tax, "NOPAT": nopat, "D&A": da, "CapEx": capex,
            "Change in NWC": dnwc, "FCFF": nopat + da - capex - dnwc, "EBITDA": ebit + da,
        }
        rev_prev = rev
    return pd.DataFrame(cols)


def _finish_dcf(fc, wacc, tv, co):
    n = np.arange(1, FORECAST_YEARS + 1)
    factors = 1 / (1 + wacc) ** n
    pv_fcff = fc.loc["FCFF"].values * factors
    pv_tv = tv * factors[-1]
    ev = pv_fcff.sum() + pv_tv
    equity = ev - co["net_debt"]
    price = max(equity / co["shares"], 0.0) * co["fx"]
    return {"factors": factors, "pv_fcff": pv_fcff, "sum_pv": pv_fcff.sum(), "tv": tv, "pv_tv": pv_tv,
            "ev": ev, "equity": equity, "price": price, "tv_share": pv_tv / ev if ev else float("nan")}


def calculate_dcf_perpetual_growth(fc, wacc, g, co):
    fcff5 = fc.loc["FCFF"].iloc[-1]
    tv = fcff5 * (1 + g) / (wacc - g)
    return _finish_dcf(fc, wacc, tv, co)


def calculate_dcf_exit_multiple(fc, wacc, multiple, co):
    tv = fc.loc["EBITDA"].iloc[-1] * multiple
    return _finish_dcf(fc, wacc, tv, co)


def calculate_sensitivity(a, co, exit_mult):
    """DCF price for each WACC x terminal growth pair (revenue growth follows g, as in the main model)."""
    perp = pd.DataFrame(index=SENS_WACC, columns=SENS_G, dtype=float)
    blend = perp.copy()
    for g in SENS_G:
        fc = project_fcff(a, growth=g)
        for w in SENS_WACC:
            p = calculate_dcf_perpetual_growth(fc, w, g, co)["price"]
            x = calculate_dcf_exit_multiple(fc, w, exit_mult, co)["price"]
            perp.loc[w, g], blend.loc[w, g] = p, (p + x) / 2
    return perp, blend


# --------------------------------------------------------------------------
# 4. Comparable company analysis
# --------------------------------------------------------------------------
def find_peer_tickers(co, user_peers=None, want=5):
    if user_peers:
        return [p for p in user_peers if p != co["ticker"].upper()][:8]
    cands = []
    for kind, key in (("Industry", co["industry_key"]), ("Sector", co["sector_key"])):
        if not key:
            continue
        top = safe(lambda: getattr(yf, kind)(key).top_companies)
        if top is not None and len(top):
            cands += [s for s in top.index.tolist() if s not in cands]
    return [c for c in cands if c.upper() != co["ticker"].upper()][:16]


def get_peer_data(tickers, target_mc=None, want=5):
    peers = []
    for s in tickers:
        info = safe(lambda: yf.Ticker(s).info, {}) or {}
        mc = num(info.get("marketCap"))
        if not ok(mc):
            continue
        ev, ebitda, pe = num(info.get("enterpriseValue")), num(info.get("ebitda")), num(info.get("trailingPE"))
        peers.append({
            "name": info.get("shortName") or s, "ticker": s, "market_cap": mc, "ev": ev,
            "revenue": num(info.get("totalRevenue")), "ebitda": ebitda,
            "pe": pe if ok(pe) and 0 < pe <= 100 else float("nan"),
            "ev_ebitda": ev / ebitda if ok(ev) and ok(ebitda) and 0 < ev / ebitda <= 50 else float("nan"),
        })
    if target_mc and len(peers) > want:  # keep the peers closest in size to the target
        peers.sort(key=lambda p: abs(math.log(p["market_cap"] / target_mc)))
        peers = peers[:want]
    return peers[:want]


def calculate_cca(co, peers):
    """Peer statistics and implied share prices. Any unusable method is skipped with an explanation."""
    def stats(key):
        v = np.array([p[key] for p in peers if ok(p[key]) and p[key] > 0])
        if len(v) == 0:
            return {"p25": float("nan"), "median": float("nan"), "p75": float("nan"), "n": 0}
        return {"p25": np.percentile(v, 25), "median": np.median(v), "p75": np.percentile(v, 75), "n": len(v)}

    pe_stats, ev_stats = stats("pe"), stats("ev_ebitda")
    notes = []
    tgt_ebitda, eps, fx = co["ebitda"], co["eps"], co["fx"]  # eps is already in price currency

    ev_price = pe_price = float("nan")
    ev_imp = eq_imp = float("nan")
    if not peers:
        notes.append("No comparable companies could be retrieved, so CCA is not available.")
    if ok(tgt_ebitda) and tgt_ebitda > 0 and ev_stats["n"]:
        ev_imp = tgt_ebitda * ev_stats["median"]
        eq_imp = ev_imp - co["net_debt"]
        ev_price = max(eq_imp / co["shares"], 0.0) * fx
    elif peers:
        why = "target EBITDA is negative or unavailable" if not (ok(tgt_ebitda) and tgt_ebitda > 0) else "no peer has a positive EV/EBITDA"
        notes.append(f"EV/EBITDA method skipped: {why}.")
    if ok(eps) and eps > 0 and pe_stats["n"]:
        pe_price = eps * pe_stats["median"]
    elif peers:
        why = "target EPS is negative or unavailable" if not (ok(eps) and eps > 0) else "no peer has a positive P/E"
        notes.append(f"P/E method skipped: {why}.")

    prices = [p for p in (ev_price, pe_price) if ok(p)]
    tgt_pe = co["price"] / eps if ok(eps) and eps > 0 else float("nan")
    tgt_evm = co["ev"] / (tgt_ebitda * fx) if ok(tgt_ebitda) and tgt_ebitda > 0 else float("nan")
    return {
        "pe": pe_stats, "ev": ev_stats, "ev_price": ev_price, "pe_price": pe_price,
        "implied_ev": ev_imp, "implied_equity": eq_imp,
        "price": float(np.mean(prices)) if prices else float("nan"),
        "target_pe": tgt_pe, "target_ev_ebitda": tgt_evm, "notes": notes,
    }


# --------------------------------------------------------------------------
# 5. Final valuation
# --------------------------------------------------------------------------
def combine_valuation(co, dcf_price, cca_price, w_dcf, w_cca):
    notes = []
    if ok(cca_price):
        final = dcf_price * w_dcf + cca_price * w_cca
        eff = (w_dcf, w_cca)
    else:
        final, eff = dcf_price, (1.0, 0.0)
        notes.append("CCA was unavailable, so the final value uses 100% DCF.")
    upside = final / co["price"] - 1
    return {"final": final, "upside": upside, "rec": "BUY" if upside >= MIN_UPSIDE else "SELL",
            "weights": eff, "notes": notes}


# --------------------------------------------------------------------------
# 6. Excel styling helpers
# --------------------------------------------------------------------------
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def na(v):
    if v is None:
        return "N/A"
    if isinstance(v, (np.floating, float)):
        return float(v) if math.isfinite(v) else "N/A"
    if isinstance(v, np.integer):
        return int(v)
    return v


def put(ws, row, col, value, fmt=None, bold=False, color="000000", fill=None, align=None,
        size=10, italic=False, wrap=False, border=False, valign="center"):
    c = ws.cell(row=row, column=col)
    c.value = None if value is None else na(value)
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    if fmt and not isinstance(c.value, str):
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    h = align or ("right" if isinstance(c.value, (int, float)) else "left")
    if c.value == "N/A" and not align:
        h = "right"
    c.alignment = Alignment(horizontal=h, vertical=valign, wrap_text=wrap)
    if border:
        c.border = BORDER
    return c


def box(ws, rng, value, **style):
    """Merge a range, styling every cell in it (so fill and borders look continuous)."""
    fill_border = {k: v for k, v in style.items() if k in ("fill", "border")}
    for row in ws[rng]:
        for c in row:
            put(ws, c.row, c.column, None, **fill_border)
    ws.merge_cells(rng)
    first = ws[rng.split(":")[0]]
    put(ws, first.row, first.column, value, **style)


def section(ws, row, text, c1=2, c2=8):
    for c in range(c1, c2 + 1):
        put(ws, row, c, text if c == c1 else None, bold=True, color="FFFFFF", fill=NAVY, size=11)
    ws.row_dimensions[row].height = 20


def header(ws, row, labels, c1=2, fill=LIGHT):
    for i, lab in enumerate(labels):
        put(ws, row, c1 + i, lab, bold=True, fill=fill, align="left" if i == 0 else "center", border=True)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def sym(co):
    return CURRENCY_SYMBOLS.get(co["currency"], co["currency"] + " ")


def px_fmt(co):
    return f'"{sym(co)}"#,##0.00'


def m(x):
    """Raw currency units -> millions for display."""
    return x / 1e6 if ok(x) else float("nan")


def title_block(ws, co, text, last_col="H"):
    box(ws, f"B1:{last_col}1", text, bold=True, size=16, color="FFFFFF", fill=NAVY)
    ws.row_dimensions[1].height = 30
    box(ws, f"B2:{last_col}2", f"{co['name']} ({co['ticker']})  |  Analysis date {dt.date.today():%d %b %Y}",
        italic=True, color="595959", size=10)


# --------------------------------------------------------------------------
# 7. Sheet builders
# --------------------------------------------------------------------------
def create_financials_sheet(wb, co):
    ws = wb.create_sheet("Financials")
    h = co["hist"]
    n = len(h)
    last_col = get_column_letter(2 + n)
    title_block(ws, co, "HISTORICAL FINANCIALS", last_col)
    box(ws, f"B3:{last_col}3", f"Source: Yahoo Finance annual statements (Historical Data). Millions of {co['fin_currency']}.",
        italic=True, color="595959", size=9)
    header(ws, 5, ["Metric"] + [f"FY{y}" for y in h.index])
    rows = [("Revenue", "Revenue", FMT_M, True), ("Revenue Growth", "Rev Growth", FMT_PCT, False),
            ("Gross Profit", "Gross Profit", FMT_M, False), ("Gross Margin", "Gross Margin", FMT_PCT, False),
            ("Operating Income", "EBIT", FMT_M, True), ("Operating Margin", "EBIT Margin", FMT_PCT, False),
            ("Net Income", "Net Income", FMT_M, True), ("Net Margin", "Net Margin", FMT_PCT, False),
            ("EBITDA", "EBITDA", FMT_M, False), ("Capital Expenditure", "CapEx", FMT_M, False),
            ("Free Cash Flow", "FCF", FMT_M, True), ("Total Debt", "Total Debt", FMT_M, False),
            ("Debt / Equity", "Debt/Equity", FMT_X.replace("0.0", "0.00"), False)]
    pos = {}
    for i, (label, col, fmt, bold) in enumerate(rows):
        r = 6 + i
        pos[col] = r
        put(ws, r, 2, label, bold=bold, border=True)
        for j, y in enumerate(h.index):
            v = h.loc[y, col]
            put(ws, r, 3 + j, m(v) if fmt == FMT_M else v, fmt=fmt, bold=bold, border=True)
    ws.freeze_panes = "C6"
    widths(ws, {"A": 2, "B": 26, **{get_column_letter(3 + j): 14 for j in range(n)}})

    # Price history table (used by the chart), far to the right
    pc = 3 + n + 2
    put(ws, 5, pc, "Date", bold=True, fill=LIGHT, border=True, align="center")
    put(ws, 5, pc + 1, f"Close ({co['currency']})", bold=True, fill=LIGHT, border=True, align="center")
    prices = co["prices"]
    for i, (d, v) in enumerate(prices.items()):
        put(ws, 6 + i, pc, d.strftime("%Y-%m-%d"), align="center")
        put(ws, 6 + i, pc + 1, float(v), fmt=FMT_NUM)
    widths(ws, {get_column_letter(pc): 12, get_column_letter(pc + 1): 14})

    cats = Reference(ws, min_col=3, max_col=2 + n, min_row=5)
    # Revenue chart
    ch = BarChart()
    ch.type, ch.title, ch.style = "col", f"Revenue (millions {co['fin_currency']})", 10
    ch.add_data(Reference(ws, min_col=2, max_col=2 + n, min_row=pos["Revenue"]), from_rows=True, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.y_axis.number_format = "#,##0"
    ch.height, ch.width = 7.5, 13
    ws.add_chart(ch, "B21")
    # Margin chart
    lc = LineChart()
    lc.title, lc.style = "Margins", 12
    for key in ("Gross Margin", "EBIT Margin", "Net Margin"):
        lc.add_data(Reference(ws, min_col=2, max_col=2 + n, min_row=pos[key]), from_rows=True, titles_from_data=True)
    lc.set_categories(cats)
    lc.x_axis.delete = lc.y_axis.delete = False
    lc.y_axis.number_format = "0%"
    lc.height, lc.width = 7.5, 13
    ws.add_chart(lc, "B37")
    # Price chart
    if len(prices):
        pcht = LineChart()
        pcht.title, pcht.style = f"Share Price - Last 12 Months ({co['currency']})", 12
        pcht.add_data(Reference(ws, min_col=pc + 1, min_row=5, max_row=5 + len(prices)), titles_from_data=True)
        pcht.set_categories(Reference(ws, min_col=pc, min_row=6, max_row=5 + len(prices)))
        pcht.legend = None
        pcht.x_axis.delete = pcht.y_axis.delete = False
        pcht.x_axis.tickLblSkip = 30
        pcht.series[0].smooth = False
        pcht.series[0].graphicalProperties.line.width = 15000
        pcht.height, pcht.width = 7.5, 15
        ws.add_chart(pcht, "F21" if n >= 3 else "F37")
    if h["Revenue"].notna().sum() < 3:
        put(ws, 20, 2, "Note: fewer than 3 years of history are available for this ticker.", italic=True, color=RED_FG)
    return ws


def create_dcf_sheet(wb, co, wacc_info, a, tags, fc, perp, exitm, wacc_used, g, mult):
    ws = wb.create_sheet("DCF")
    widths(ws, {"A": 2, "B": 40, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 60})
    title_block(ws, co, "DISCOUNTED CASH FLOW (DCF) VALUATION")
    box(ws, "B3:H3", f"Millions of {co['fin_currency']} unless per share.  Blue = assumption / user input.  "
        f"Every figure is labelled Historical Data, Projected, or Assumption.", italic=True, color="595959", size=9)
    pxf = px_fmt(co)
    r = 5

    # 1. Historical
    h = co["hist"]
    section(ws, r, "1. HISTORICAL DATA (from Yahoo Finance)")
    r += 1
    header(ws, r, ["(millions)"] + [f"FY{y}" for y in h.index])
    r += 1
    for label, col, fmt in [("Revenue", "Revenue", FMT_M), ("Revenue Growth", "Rev Growth", FMT_PCT),
                            ("EBIT / Operating Income", "EBIT", FMT_M), ("EBIT Margin  (EBIT / Revenue)", "EBIT Margin", FMT_PCT),
                            ("Depreciation & Amortisation", "D&A", FMT_M), ("Capital Expenditure", "CapEx", FMT_M),
                            ("Net Working Capital (ex cash & ST debt)", "NWC", FMT_M), ("Cash", "Cash", FMT_M),
                            ("Total Debt", "Total Debt", FMT_M), ("Shares Outstanding (m)", "Shares", FMT_M)]:
        put(ws, r, 2, label, border=True)
        for j, y in enumerate(h.index):
            v = h.loc[y, col]
            put(ws, r, 3 + j, m(v) if fmt == FMT_M else v, fmt=fmt, border=True)
        r += 1
    put(ws, r, 2, "Net Debt", bold=True, border=True)
    for j, y in enumerate(h.index):
        put(ws, r, 3 + j, m(h.loc[y, "Total Debt"] - h.loc[y, "Cash"]), fmt=FMT_M, bold=True, border=True)
    r += 2

    # 2. WACC
    section(ws, r, "2. WACC  (calculated reference)")
    r += 1
    header(ws, r, ["Input", "Value", "Source"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    r += 1
    wrows = [("Risk-Free Rate", "rf", FMT_PCT), ("Beta", "beta", "0.00"), ("Equity Risk Premium", "erp", FMT_PCT),
             ("Cost of Equity = Rf + Beta x ERP", "coe", FMT_PCT), ("Pre-Tax Cost of Debt", "kd", FMT_PCT),
             ("Tax Rate", "tax", FMT_PCT), (f"Market Value of Equity (m {co['currency']})", "E", FMT_M),
             (f"Market Value of Debt (m {co['currency']})", "D", FMT_M),
             ("Equity Weight  E/V", "we", FMT_PCT), ("Debt Weight  D/V", "wd", FMT_PCT)]
    for label, key, fmt in wrows:
        val, src = wacc_info[key]
        put(ws, r, 2, label, border=True)
        put(ws, r, 3, m(val) if fmt == FMT_M else val, fmt=fmt, border=True)
        put(ws, r, 4, src, italic=True, color="595959", size=9)
        r += 1
    put(ws, r, 2, "Calculated WACC = E/V x Ke + D/V x Kd x (1 - t)", bold=True, border=True, fill=GREY)
    put(ws, r, 3, wacc_info["wacc"], fmt=FMT_PCT, bold=True, border=True, fill=GREY)
    r += 1
    put(ws, r, 2, "WACC USED IN DCF (user selection)", bold=True, border=True, fill=LIGHT)
    put(ws, r, 3, wacc_used, fmt=FMT_PCT, bold=True, border=True, fill=LIGHT, color=INPUT_BLUE)
    put(ws, r, 4, "User assumption (default 9.0%, allowed 7%-12%)", italic=True, color="595959", size=9)
    r += 2

    # 3. Forecast assumptions
    section(ws, r, "3. FORECAST ASSUMPTIONS")
    r += 1
    header(ws, r, ["Driver", "Value", "Source"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    r += 1
    for label, key, fmt in [("Base-Year Revenue (m)", "base_revenue", FMT_M), ("Revenue Growth (each year)", "growth", FMT_PCT),
                            ("EBIT Margin", "ebit_margin", FMT_PCT), ("D&A % of Revenue", "da_pct", FMT_PCT),
                            ("CapEx % of Revenue", "capex_pct", FMT_PCT), ("NWC % of Revenue", "nwc_pct", FMT_PCT),
                            ("Tax Rate", "tax", FMT_PCT)]:
        v = a[key]
        put(ws, r, 2, label, border=True)
        is_assump = tags[key].startswith(("User", "Assumption"))
        put(ws, r, 3, m(v) if fmt == FMT_M else v, fmt=fmt, border=True, color=INPUT_BLUE if is_assump else "000000")
        put(ws, r, 4, tags[key], italic=True, color="595959", size=9)
        r += 1
    r += 1

    # 4. Projection
    section(ws, r, "4. FIVE-YEAR FCFF PROJECTION  (all figures Projected)")
    r += 1
    header(ws, r, ["(millions)"] + list(fc.columns))
    r += 1
    put(ws, r, 2, "Status", italic=True, border=True)
    for j in range(FORECAST_YEARS):
        put(ws, r, 3 + j, "Projected", italic=True, align="center", border=True, color="7F7F7F")
    r += 1
    proj_rows = [("Revenue", "Revenue", FMT_M, True), ("Revenue Growth", "Revenue Growth", FMT_PCT, False),
                 ("EBIT", "EBIT", FMT_M, False), ("EBIT Margin", "EBIT Margin", FMT_PCT, False),
                 ("Less: Taxes on EBIT", "Taxes on EBIT", FMT_M, False), ("NOPAT = EBIT x (1 - t)", "NOPAT", FMT_M, True),
                 ("Plus: D&A", "D&A", FMT_M, False), ("Less: Capital Expenditure", "CapEx", FMT_M, False),
                 ("Less: Change in NWC", "Change in NWC", FMT_M, False), ("FCFF = NOPAT + D&A - CapEx - dNWC", "FCFF", FMT_M, True),
                 ("Memo: EBITDA (EBIT + D&A)", "EBITDA", FMT_M, False)]
    for label, key, fmt, bold in proj_rows:
        put(ws, r, 2, label, bold=bold, border=True, fill=GREY if key == "FCFF" else None)
        for j in range(FORECAST_YEARS):
            v = fc.loc[key].iloc[j]
            put(ws, r, 3 + j, m(v) if fmt == FMT_M else v, fmt=fmt, bold=bold, border=True, fill=GREY if key == "FCFF" else None)
        r += 1
    put(ws, r, 2, "Discount Factor  1/(1+WACC)^t", border=True)
    for j in range(FORECAST_YEARS):
        put(ws, r, 3 + j, perp["factors"][j], fmt="0.0000", border=True)
    r += 1
    put(ws, r, 2, "PV of FCFF", bold=True, border=True)
    for j in range(FORECAST_YEARS):
        put(ws, r, 3 + j, m(perp["pv_fcff"][j]), fmt=FMT_M, bold=True, border=True)
    r += 2

    # 5-6. Two DCF methods side by side
    section(ws, r, "5. DCF VALUATION - TWO TERMINAL VALUE METHODS")
    r += 1
    header(ws, r, ["(millions)", "Perpetual Growth", "Exit Multiple"])
    r += 1
    fcff5 = fc.loc["FCFF"].iloc[-1]
    lines = [
        ("Terminal assumption", f"g = {g:.1%}", f"{mult:.1f}x EBITDA", None, False),
        ("Sum of PV of Forecast FCFF", m(perp["sum_pv"]), m(exitm["sum_pv"]), FMT_M, False),
        ("Terminal Value (undiscounted)", m(perp["tv"]), m(exitm["tv"]), FMT_M, False),
        ("PV of Terminal Value", m(perp["pv_tv"]), m(exitm["pv_tv"]), FMT_M, False),
        ("Enterprise Value", m(perp["ev"]), m(exitm["ev"]), FMT_M, True),
        ("Less: Net Debt", m(co["net_debt"]), m(co["net_debt"]), FMT_M, False),
        ("Equity Value", m(perp["equity"]), m(exitm["equity"]), FMT_M, True),
        ("Shares Outstanding (m)", m(co["shares"]), m(co["shares"]), FMT_M, False),
    ]
    if co["fx"] != 1:
        lines.append((f"FX rate ({co['fin_currency']} -> {co['currency']})", co["fx"], co["fx"], "0.0000", False))
    lines += [("Implied Share Price", perp["price"], exitm["price"], pxf, True),
              ("Terminal Value as % of EV", perp["tv_share"], exitm["tv_share"], FMT_PCT, False)]
    for label, v1, v2, fmt, bold in lines:
        put(ws, r, 2, label, bold=bold, border=True, fill=GREY if label == "Implied Share Price" else None)
        for j, v in enumerate((v1, v2)):
            put(ws, r, 3 + j, v, fmt=fmt, bold=bold, border=True, align="right", fill=GREY if label == "Implied Share Price" else None)
        r += 1
    put(ws, r - 3 if co["fx"] == 1 else r - 4, 8, f"TV(perpetual) = FCFF5 x (1+g) / (WACC - g);  TV(exit) = EBITDA5 x multiple", italic=True, color="595959", size=9)
    r += 1
    put(ws, r, 2, "DCF IMPLIED PRICE (average of the two methods)", bold=True, fill=LIGHT, border=True)
    put(ws, r, 3, (perp["price"] + exitm["price"]) / 2, fmt=pxf, bold=True, fill=LIGHT, border=True)
    ws.freeze_panes = "A5"
    if perp["ev"] <= 0 or fcff5 <= 0:
        put(ws, r + 2, 2, "WARNING: Year-5 FCFF or enterprise value is not positive - the DCF is unreliable for this company.",
            bold=True, color=RED_FG)
    return ws


def create_cca_sheet(wb, co, peers, cca):
    ws = wb.create_sheet("CCA")
    widths(ws, {"A": 2, "B": 34, "C": 12, "D": 16, "E": 16, "F": 16, "G": 16, "H": 12, "I": 14})
    title_block(ws, co, "COMPARABLE COMPANY ANALYSIS (CCA)", "I")
    box(ws, "B3:I3", "Peers from Yahoo Finance industry/sector lists (or user-supplied). Market data, millions in each company's own reporting/trading currency. "
        "Multiples that are negative, undefined or extreme (P/E > 100x, EV/EBITDA > 50x) are shown as N/A and excluded from the statistics.", italic=True, color="595959", size=9, wrap=True)
    ws.row_dimensions[3].height = 26
    section(ws, 5, "PEER GROUP", 2, 9)
    header(ws, 6, ["Company", "Ticker", "Market Cap (m)", "Enterprise Value (m)", "Revenue (m)", "EBITDA (m)", "P/E", "EV/EBITDA"])
    tgt = {"name": co["name"], "ticker": co["ticker"], "market_cap": co["market_cap"], "ev": co["ev"],
           "revenue": co["hist"]["Revenue"].dropna().iloc[-1] * co["fx"], "ebitda": co["ebitda"] * co["fx"] if ok(co["ebitda"]) else float("nan"),
           "pe": cca["target_pe"], "ev_ebitda": cca["target_ev_ebitda"]}
    r = 7
    for i, p in enumerate([tgt] + peers):
        fill = LIGHT if i == 0 else None
        vals = [p["name"] + (" (Target)" if i == 0 else ""), p["ticker"], m(p["market_cap"]), m(p["ev"]), m(p["revenue"]),
                m(p["ebitda"]), p["pe"], p["ev_ebitda"]]
        fmts = [None, None, FMT_M, FMT_M, FMT_M, FMT_M, FMT_X, FMT_X]
        for j, (v, f) in enumerate(zip(vals, fmts)):
            put(ws, r, 2 + j, v, fmt=f, border=True, fill=fill, bold=(i == 0))
        r += 1
    if not peers:
        put(ws, r, 2, "No comparable companies could be retrieved.", italic=True, color=RED_FG)
        r += 1
    r += 1
    section(ws, r, "PEER STATISTICS", 2, 9)
    r += 1
    header(ws, r, ["Statistic", "", "", "", "", "", "P/E", "EV/EBITDA"])
    r += 1
    stat_rows = [("25th Percentile", "p25"), ("Median", "median"), ("75th Percentile", "p75")]
    stat_start = r
    for label, key in stat_rows:
        put(ws, r, 2, label, border=True, bold=(key == "median"))
        for c in range(3, 8):
            put(ws, r, c, None, border=True)
        put(ws, r, 8, cca["pe"][key], fmt=FMT_X, border=True, bold=(key == "median"))
        put(ws, r, 9, cca["ev"][key], fmt=FMT_X, border=True, bold=(key == "median"))
        r += 1
    put(ws, r, 2, "Number of valid peers", border=True, italic=True)
    for c in range(3, 8):
        put(ws, r, c, None, border=True)
    put(ws, r, 8, cca["pe"]["n"], border=True, fmt="0")
    put(ws, r, 9, cca["ev"]["n"], border=True, fmt="0")
    r += 2

    section(ws, r, "IMPLIED VALUATION OF TARGET", 2, 9)
    r += 1
    pxf = px_fmt(co)
    put(ws, r, 2, "A. EV/EBITDA method", bold=True)
    r += 1
    items = [("Target EBITDA (m)", m(co["ebitda"] * co["fx"]) if ok(co["ebitda"]) else float("nan"), FMT_M),
             ("x Median EV/EBITDA", cca["ev"]["median"], FMT_X),
             ("= Implied Enterprise Value (m)", m(cca["implied_ev"] * co["fx"]) if ok(cca["implied_ev"]) else float("nan"), FMT_M),
             ("- Net Debt (m)", m(co["net_debt"] * co["fx"]), FMT_M),
             ("= Implied Equity Value (m)", m(cca["implied_equity"] * co["fx"]) if ok(cca["implied_equity"]) else float("nan"), FMT_M),
             ("/ Shares Outstanding (m)", m(co["shares"]), FMT_M),
             ("= EV/EBITDA Implied Share Price", cca["ev_price"], pxf)]
    for label, v, f in items:
        put(ws, r, 2, label, border=True, bold=label.startswith("= EV"))
        put(ws, r, 3, v, fmt=f, border=True, bold=label.startswith("= EV"))
        r += 1
    r += 1
    put(ws, r, 2, "B. P/E method", bold=True)
    r += 1
    for label, v, f in [("Target EPS (trailing)", co["eps"] if ok(co["eps"]) else float("nan"), pxf),
                        ("x Median P/E", cca["pe"]["median"], FMT_X), ("= P/E Implied Share Price", cca["pe_price"], pxf)]:
        put(ws, r, 2, label, border=True, bold=label.startswith("= P/E"))
        put(ws, r, 3, v, fmt=f, border=True, bold=label.startswith("= P/E"))
        r += 1
    r += 1
    put(ws, r, 2, "CCA IMPLIED PRICE (average of available methods)", bold=True, fill=LIGHT, border=True)
    put(ws, r, 3, cca["price"], fmt=pxf, bold=True, fill=LIGHT, border=True)
    r += 1
    for n in cca["notes"]:
        put(ws, r, 2, "Note: " + n, italic=True, color=RED_FG)
        r += 1
    r += 1

    # Chart data
    section(ws, r, "CHART DATA", 2, 9)
    r += 1
    header(ws, r, ["", "P/E", "EV/EBITDA"])
    cstart = r
    for label, pe, ev in [("Target", cca["target_pe"], cca["target_ev_ebitda"]), ("25th Percentile", cca["pe"]["p25"], cca["ev"]["p25"]),
                          ("Median", cca["pe"]["median"], cca["ev"]["median"]), ("75th Percentile", cca["pe"]["p75"], cca["ev"]["p75"])]:
        r += 1
        put(ws, r, 2, label, border=True)
        put(ws, r, 3, pe, fmt=FMT_X, border=True)
        put(ws, r, 4, ev, fmt=FMT_X, border=True)
    for i, (title, col, anchor) in enumerate([("EV/EBITDA: Target vs Peers", 4, "B" + str(r + 2)), ("P/E: Target vs Peers", 3, "F" + str(r + 2))]):
        ch = BarChart()
        ch.type, ch.title, ch.style = "col", title, 10
        ch.add_data(Reference(ws, min_col=col, min_row=cstart, max_row=cstart + 4), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=cstart + 1, max_row=cstart + 4))
        ch.legend = None
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.height, ch.width = 7.5, 11
        for k, colr in enumerate([BLUE, "A6A6A6", "7F7F7F", "A6A6A6"]):
            dp = DataPoint(idx=k)
            dp.graphicalProperties.solidFill = colr
            ch.series[0].dPt.append(dp)
        ws.add_chart(ch, anchor)
    ws.freeze_panes = "A5"
    return ws


def create_sensitivity_sheet(wb, co, perp_grid, blend_grid, g, wacc_used):
    ws = wb.create_sheet("Sensitivity")
    widths(ws, {"A": 2, "B": 18, **{get_column_letter(c): 12 for c in range(3, 8)}})
    title_block(ws, co, "DCF SENSITIVITY ANALYSIS", "G")
    box(ws, "B3:G3", "Implied share price for each WACC / terminal growth pair. Revenue growth follows the terminal growth rate, "
        "exactly as in the main model. Green = higher value, red = lower.", italic=True, color="595959", size=9, wrap=True)
    ws.row_dimensions[3].height = 26
    r = 5
    for title, grid in [("1. Perpetual Growth DCF - Implied Share Price", perp_grid),
                        ("2. Blended DCF (average of Perpetual Growth & Exit Multiple) - Implied Share Price", blend_grid)]:
        section(ws, r, title, 2, 7)
        r += 1
        put(ws, r, 2, "WACC \\ Growth", bold=True, fill=LIGHT, border=True, align="center")
        for j, gg in enumerate(SENS_G):
            put(ws, r, 3 + j, gg, fmt=FMT_PCT, bold=True, fill=LIGHT, border=True, align="center")
        first = r + 1
        for w in SENS_WACC:
            r += 1
            put(ws, r, 2, w, fmt=FMT_PCT, bold=True, fill=LIGHT, border=True, align="center")
            for j, gg in enumerate(SENS_G):
                put(ws, r, 3 + j, grid.loc[w, gg], fmt=px_fmt(co), border=True, align="center")
        ws.conditional_formatting.add(f"C{first}:G{r}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B"))
        r += 2
    put(ws, r, 2, f"Current share price: {sym(co)}{co['price']:,.2f}   |   Base case used in model: WACC {wacc_used:.1%}, growth {g:.1%}",
        bold=True)
    return ws


def create_dashboard(wb, co, res, dcf_price, cca_price, perp_price, exit_price, wacc_used, g, mult, w_dcf, w_cca, notes):
    ws = wb.active
    ws.title = "Dashboard"
    widths(ws, {"A": 2, **{get_column_letter(c): 15 for c in range(2, 10)}, "J": 2})
    pxf = px_fmt(co)
    box(ws, "B1:I1", "COMPANY FINANCIAL VALUATION", bold=True, size=20, color="FFFFFF", fill=NAVY, align="center")
    ws.row_dimensions[1].height = 38
    box(ws, "B2:I2", f"{co['name']}  ({co['ticker']})", bold=True, size=14, color=NAVY, fill=LIGHT, align="center")
    ws.row_dimensions[2].height = 26

    info = [("Company", co["name"], "Analysis Date", f"{dt.date.today():%d %b %Y}"),
            ("Ticker", co["ticker"], "Current Share Price", (co["price"], pxf)),
            ("Sector", co["sector"], f"Market Cap (m {co['currency']})", (m(co["market_cap"]), FMT_M)),
            ("Industry", co["industry"], f"Enterprise Value (m {co['currency']})", (m(co["ev"]), FMT_M))]
    for i, (l1, v1, l2, v2) in enumerate(info):
        r = 4 + i
        put(ws, r, 2, l1, bold=True, color="595959")
        box(ws, f"C{r}:E{r}", v1, border=False)
        put(ws, r, 6, l2, bold=True, color="595959")
        v, f = (v2 if isinstance(v2, tuple) else (v2, None))
        box(ws, f"G{r}:I{r}", v, fmt=f, align="left")

    # Valuation cards
    cards = [("B", "C", "CURRENT PRICE", co["price"], LIGHT), ("D", "E", "DCF IMPLIED PRICE", dcf_price, LIGHT),
             ("F", "G", "CCA IMPLIED PRICE", cca_price, LIGHT), ("H", "I", "FINAL INTRINSIC PRICE", res["final"], "FFF2CC")]
    for c1, c2, label, v, fill in cards:
        box(ws, f"{c1}9:{c2}9", label, bold=True, size=10, color="FFFFFF", fill=BLUE, align="center", border=True)
        box(ws, f"{c1}10:{c2}11", v, fmt=pxf, bold=True, size=20, color=NAVY, fill=fill, align="center", border=True)
    ws.row_dimensions[10].height = ws.row_dimensions[11].height = 20

    up = res["upside"]
    buy = res["rec"] == "BUY"
    box(ws, "B13:E13", "UPSIDE / DOWNSIDE", bold=True, color="FFFFFF", fill=NAVY, align="center", border=True)
    box(ws, "F13:I13", "RECOMMENDATION", bold=True, color="FFFFFF", fill=NAVY, align="center", border=True)
    box(ws, "B14:E16", up, fmt='+0.0%;-0.0%;0.0%', bold=True, size=28, color=GREEN_FG if buy else RED_FG,
        fill=GREEN_BG if buy else RED_BG, align="center", border=True)
    box(ws, "F14:I16", res["rec"], bold=True, size=28, color=GREEN_FG if buy else RED_FG,
        fill=GREEN_BG if buy else RED_BG, align="center", border=True)
    box(ws, "B17:I17", f"Rule: BUY if upside >= {MIN_UPSIDE:.0%}, otherwise SELL.  Upside = Final Intrinsic Price / Current Price - 1.",
        italic=True, size=9, color="595959", align="center")
    box(ws, "B18:I18", "Recommendation is based solely on the model's valuation assumptions and minimum required upside threshold. "
        "It is not investment advice.", italic=True, size=9, color="595959", align="center", wrap=True)
    ws.row_dimensions[18].height = 26

    # Chart data + assumptions
    section(ws, 20, "VALUATION COMPARISON", 2, 9)
    header(ws, 21, ["Measure", f"Price ({co['currency']})"])
    for i, (lab, v) in enumerate([("Current Price", co["price"]), ("DCF Value", dcf_price), ("CCA Value", cca_price), ("Final Intrinsic", res["final"])]):
        put(ws, 22 + i, 2, lab, border=True)
        put(ws, 22 + i, 3, v, fmt=pxf, border=True, bold=(i == 3))
    put(ws, 27, 2, "KEY ASSUMPTIONS", bold=True, color="FFFFFF", fill=NAVY)
    put(ws, 27, 3, None, fill=NAVY)
    for i, (lab, v, f) in enumerate([("WACC (used)", wacc_used, FMT_PCT), ("Terminal Growth", g, FMT_PCT), ("Exit Multiple", mult, FMT_X),
                                     ("DCF Weight", res["weights"][0], "0%"), ("CCA Weight", res["weights"][1], "0%"),
                                     ("Perpetual Growth DCF", perp_price, pxf), ("Exit Multiple DCF", exit_price, pxf)]):
        put(ws, 28 + i, 2, lab, border=True)
        put(ws, 28 + i, 3, v, fmt=f, border=True, color=INPUT_BLUE if i < 5 else "000000")

    ch = BarChart()
    ch.type, ch.title, ch.style = "col", "Current Price vs Valuation", 10
    ch.add_data(Reference(ws, min_col=3, min_row=21, max_row=25), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=22, max_row=25))
    ch.legend = None
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.y_axis.number_format = "#,##0"
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    for k, colr in enumerate(["7F7F7F", BLUE, "5B9BD5", "1F3864"]):
        dp = DataPoint(idx=k)
        dp.graphicalProperties.solidFill = colr
        ch.series[0].dPt.append(dp)
    ch.height, ch.width = 9.5, 14.5
    ws.add_chart(ch, "D20")

    r = 37
    section(ws, r, "DATA NOTES & ASSUMPTIONS", 2, 9)
    lines = notes or ["No data issues were detected; all inputs come from Yahoo Finance or the stated user assumptions."]
    for i, t in enumerate(lines):
        box(ws, f"B{r + 1 + i}:I{r + 1 + i}", "- " + t, size=9, wrap=True, color="404040")
        ws.row_dimensions[r + 1 + i].height = 26 if len(t) > 110 else 15
    ws.sheet_view.showGridLines = False
    return ws


def format_workbook(wb):
    """Order sheets, hide gridlines, and set print / tab settings."""
    order = ["Dashboard", "DCF", "CCA", "Financials", "Sensitivity"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    colors = {"Dashboard": NAVY, "DCF": BLUE, "CCA": "5B9BD5", "Financials": "7F7F7F", "Sensitivity": "A6A6A6"}
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = colors.get(ws.title, BLUE)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0


def save_workbook(wb, ticker):
    name = f"{ticker.upper().replace('^', '')}_Financial_Valuation.xlsx"
    path = os.path.abspath(name)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# 8. User input
# --------------------------------------------------------------------------
def ask(prompt):
    try:
        return input(prompt).strip()
    except EOFError:
        return ""


def parse_pct(text):
    """'3', '3%', '0.03' -> 0.03.  Values <= 1 without a % sign are treated as fractions."""
    t = text.replace("%", "").strip()
    v = float(t)
    return v / 100 if ("%" in text or v > 1) else v


def ask_value(label, default, lo, hi, kind="pct"):
    unit = (lambda x: f"{x:.1%}") if kind == "pct" else (lambda x: f"{x:.1f}x")
    while True:
        raw = ask(f"{label} [{unit(lo)} - {unit(hi)}, default {unit(default)}]: ")
        if not raw:
            return default
        try:
            v = parse_pct(raw) if kind == "pct" else float(raw.lower().rstrip("x"))
        except ValueError:
            print("  Invalid number - please try again.")
            continue
        if lo <= v <= hi:
            return v
        print(f"  Value must be between {unit(lo)} and {unit(hi)}.")


def ask_assumptions(calc_wacc):
    print("\nValuation assumptions - press Enter to accept each default.")
    print(f"  (Reference: WACC calculated from market data = {calc_wacc:.1%}; type 'calc' to use it.)")
    lo, hi = RANGES["wacc"]
    while True:
        raw = ask(f"Discount rate / WACC [{lo:.1%} - {hi:.1%}, default {DEFAULTS['wacc']:.1%}]: ")
        if raw.lower() == "calc":
            wacc = min(max(calc_wacc, lo), hi)
            if wacc != calc_wacc:
                print(f"  Calculated WACC is outside the allowed range; using {wacc:.1%}.")
            break
        if not raw:
            wacc = DEFAULTS["wacc"]
            break
        try:
            wacc = parse_pct(raw)
        except ValueError:
            print("  Invalid number - please try again.")
            continue
        if lo <= wacc <= hi:
            break
        print(f"  Value must be between {lo:.1%} and {hi:.1%}.")
    g = ask_value("Terminal growth rate", DEFAULTS["g"], *RANGES["g"])
    mult = ask_value("Exit EV/EBITDA multiple", DEFAULTS["exit"], *RANGES["exit"], kind="x")
    while True:
        w_dcf = ask_value("DCF weight", DEFAULTS["w_dcf"], 0.0, 1.0)
        w_cca = ask_value("CCA weight", DEFAULTS["w_cca"] if w_dcf == DEFAULTS["w_dcf"] else round(1 - w_dcf, 4), 0.0, 1.0)
        if abs(w_dcf + w_cca - 1) < 1e-6:
            break
        print(f"  DCF weight + CCA weight must equal 100% (you entered {w_dcf + w_cca:.0%}). Please re-enter both.")
    return {"wacc": wacc, "g": g, "exit": mult, "w_dcf": w_dcf, "w_cca": w_cca}


def ask_ticker():
    while True:
        t = ask("Enter company ticker: ").upper()
        if t:
            return t
        print("  Please type a ticker such as AAPL or BHP.AX.")


# --------------------------------------------------------------------------
# 9. Main
# --------------------------------------------------------------------------
def main():
    print("=" * 56)
    print("   FINANCIAL VALUATION DASHBOARD")
    print("=" * 56)
    while True:
        ticker = ask_ticker()
        print(f"\nDownloading data for {ticker} ...")
        try:
            co = get_company_data(ticker)
            break
        except DataError as e:
            print(f"  Error: {e}\n")
        except Exception as e:  # network / API failure
            print(f"  Could not retrieve data for '{ticker}' ({type(e).__name__}). Check the ticker and your connection.\n")

    print(f"Found: {co['name']} - price {sym(co)}{co['price']:,.2f} {co['currency']}")
    wacc_info = calculate_wacc(co)
    peer_in = ask("Peer tickers, comma-separated (optional, press Enter to auto-select): ")
    user_peers = [p.strip().upper() for p in peer_in.split(",") if p.strip()]
    ass = ask_assumptions(wacc_info["wacc"])

    print("\nBuilding model ...")
    a, tags = build_forecast_assumptions(co, ass["g"])
    fc = project_fcff(a)
    perp = calculate_dcf_perpetual_growth(fc, ass["wacc"], ass["g"], co)
    exitm = calculate_dcf_exit_multiple(fc, ass["wacc"], ass["exit"], co)
    dcf_price = (perp["price"] + exitm["price"]) / 2

    print("Finding comparable companies ...")
    peers = get_peer_data(find_peer_tickers(co, user_peers), co["market_cap"])
    cca = calculate_cca(co, peers)
    res = combine_valuation(co, dcf_price, cca["price"], ass["w_dcf"], ass["w_cca"])
    perp_grid, blend_grid = calculate_sensitivity(a, co, ass["exit"])

    sources = [wacc_info[k][1] for k in ("rf", "beta", "kd", "tax")] + [tags[k] for k in ("ebit_margin", "da_pct", "capex_pct", "nwc_pct", "tax")]
    notes = list(co["notes"]) + sorted({s for s in sources if s.startswith("Assumption")})
    notes += ["Equity risk premium of 5.0% is an assumption (not retrieved from data).",
              "Forecast revenue growth equals the terminal growth rate in every year (per model design); other drivers use 3-yr historical averages."]
    notes += cca["notes"] + res["notes"]
    if len(peers) < 3:
        notes.append(f"Only {len(peers)} comparable companies were found (3-5 preferred).")
    if a["ebit_margin"] <= 0 or fc.loc["FCFF"].iloc[-1] <= 0:
        notes.append("WARNING: forecast FCFF is not positive; the DCF is unreliable for this company.")
    if abs(ass["wacc"] - wacc_info["wacc"]) > 0.02:
        notes.append(f"WACC used ({ass['wacc']:.1%}) differs from the calculated WACC ({wacc_info['wacc']:.1%}).")

    wb = Workbook()
    create_dashboard(wb, co, res, dcf_price, cca["price"], perp["price"], exitm["price"], ass["wacc"], ass["g"],
                     ass["exit"], ass["w_dcf"], ass["w_cca"], notes)
    create_dcf_sheet(wb, co, wacc_info, a, tags, fc, perp, exitm, ass["wacc"], ass["g"], ass["exit"])
    create_cca_sheet(wb, co, peers, cca)
    create_financials_sheet(wb, co)
    create_sensitivity_sheet(wb, co, perp_grid, blend_grid, ass["g"], ass["wacc"])
    format_workbook(wb)
    path = save_workbook(wb, co["ticker"])

    print("\n" + "-" * 56)
    print(f"  Current price       : {sym(co)}{co['price']:,.2f}")
    print(f"  DCF implied price   : {sym(co)}{dcf_price:,.2f}")
    print(f"  CCA implied price   : {sym(co)}{cca['price']:,.2f}" if ok(cca["price"]) else "  CCA implied price   : N/A")
    print(f"  Final intrinsic     : {sym(co)}{res['final']:,.2f}")
    print(f"  Upside / downside   : {res['upside']:+.1%}")
    print(f"  Recommendation      : {res['rec']}")
    print("-" * 56)
    print(f"Saved: {path}")
    if IN_COLAB:
        from google.colab import files
        files.download(path)
    return path


if __name__ == "__main__":
    main()
